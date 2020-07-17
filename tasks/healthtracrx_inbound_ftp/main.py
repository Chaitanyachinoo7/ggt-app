import openpyxl 
import os

files = [f for f in os.listdir('./inbound/') if os.path.isfile(f)]
for f in files:
  print (f)

'''  
# Give the location of the file 
path = "C:\\Users\\Admin\\Desktop\\demo.xlsx"
  
# workbook object is created 
wb_obj = openpyxl.load_workbook(path) 
  
sheet_obj = wb_obj.active 
m_row = sheet_obj.max_row 
  
# Loop will print all values 
# of first column  
for i in range(1, m_row + 1): 
    cell_obj = sheet_obj.cell(row = i, column = 1) 
    print(cell_obj.value) 
'''

